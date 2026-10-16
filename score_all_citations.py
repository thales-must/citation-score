"""Fill missing scores for text-ready, author-valid citations to 400 fixed seeds.

From the repository root (paths do not depend on the working directory)::

    python src/score_all_citations.py                         # read-only audit
    python src/score_all_citations.py --write --limit 3        # small trial
    python src/score_all_citations.py --write --branch abstract --limit 3
    python src/score_all_citations.py --write --component author  # no LLM calls
    python src/score_all_citations.py --recompute-author      # audit full author refresh
    python src/score_all_citations.py --write --recompute-author --max-llm-calls 10000
    python src/score_all_citations.py --write --max-llm-calls 10000

The default cohort matches data_sampling.ipynb: a cited-paper abstract plus
context (or, without context, a citing-paper abstract), with valid author
lists on both papers. Author IDs, names and contiguous order are checked;
publication dates and existing scores do not determine cohort membership.
Counts are calculated afresh rather than hardcoded to 4,819.

Semantic scores are updated only when NULL. Context-bearing records use cross_score and
llm_score; context-missing records use abstract_cross_score and
abstract_llm_score. Every selected NULL author_independence is attempted in a
separate author-first pass within the eligible cohort, regardless of whether
the semantic scores already exist. Existing semantic zeros are preserved.
Audit flags are not database fields. Author relationships use the local
co-authorship snapshot built from papers whose stored publication year is
2024. No citing-date cutoff is used; these are snapshot relationships, not
verified pre-citation collaborations. --recompute-author explicitly refreshes
all eligible independence scores under this policy, journaling previous values.
Without it, existing independence scores are preserved and may reflect an older
policy. Missing author inputs remain NULL, never an assumed independent score. No alpha,
lambda, beta, normalization, or paper-level aggregation is applied here.

The existing SPECTER abstract_score is never touched. Re-running resumes from
the database and an input/model/settings-matched local result journal. The
journal preserves completed semantic results; --recompute-author deliberately
revisits the entire eligible author cohort on each run. The
journal contains private citation IDs and model explanations: do not publish
it automatically. API failures and missing inputs remain NULL, never zero.

Requires SQLAlchemy, psycopg[binary], pgvector, python-dotenv, openpyxl,
openai, sentence-transformers and torch. Uses existing .env variables without logging
credentials. Default model/settings follow the project's existing scoring
scripts; --temperature omit supports endpoints rejecting temperature=0.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
import os
import random
import sys
import time
from collections import Counter, defaultdict
from contextlib import contextmanager
from datetime import datetime, timezone
from pathlib import Path
from urllib.parse import urlsplit

from sqlalchemy import URL, create_engine, func, inspect, select, update
from sqlalchemy.orm import Session, aliased
from sqlalchemy.pool import NullPool

from postgres.model import Author, Citation, Paper, PaperAuthor

from citation_prompts import (
    ABSTRACT_PROMPT_VERSION, CONTEXT_PROMPT_VERSION, SYSTEM_PROMPT, build_prompt,
)

BASE = Path(__file__).resolve().parent
CROSS_MODEL = "cross-encoder/ms-marco-MiniLM-L-6-v2"
CROSS_REVISION = "c5ee24cb16019beea0893ab7796b1df96625c6b8"
AUTHOR_GRAPH_YEAR = 2024
AUTHOR_POLICY = "local_2024_coauthorship_snapshot_v1"
FIELDS = ("cross_score", "llm_score", "abstract_cross_score",
          "abstract_llm_score", "author_independence")

CitingPaper = aliased(Paper, name="citing_paper")
CitedPaper = aliased(Paper, name="cited_paper")


def citation_statement():
    """Use the existing mapped schema rather than duplicating column SQL."""
    return (select(
        Citation.citing_corpus_id, Citation.cited_corpus_id, Citation.contexts,
        *(getattr(Citation, name) for name in FIELDS),
        CitingPaper.abstract.label("citing_abstract"),
        CitedPaper.abstract.label("cited_abstract"),
        CitingPaper.publication_date.label("citing_date"),
        CitingPaper.year.label("citing_year"))
        .select_from(Citation)
        .join(CitingPaper, CitingPaper.corpus_id == Citation.citing_corpus_id)
        .join(CitedPaper, CitedPaper.corpus_id == Citation.cited_corpus_id))


class StopRun(RuntimeError):
    """Stop safely without exposing a credential-bearing exception."""


class MissingInput(ValueError):
    pass


def utc_now():
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def digest(value):
    return hashlib.sha256(json.dumps(value, ensure_ascii=False, sort_keys=True,
                                    default=str, allow_nan=False).encode()).hexdigest()


def clean_text(value):
    return " ".join(value.split()) if isinstance(value, str) else ""


def context_list(value):
    """Empty/whitespace contexts are unavailable; reject unknown structures."""
    if value is None:
        return []
    if isinstance(value, str):
        # JSONB is already decoded by psycopg. A literal "[1]" can be a
        # citation marker: never reinterpret an individual context as JSON.
        text = clean_text(value)
        return [text] if text else []
    if isinstance(value, dict):
        for key in ("text", "context", "sentence"):
            if key in value:
                return context_list(value[key])
        raise MissingInput("unrecognized_context_structure")
    if isinstance(value, list):
        return [text for item in value for text in context_list(item)]
    raise MissingInput("unrecognized_context_structure")


def score_value(value):
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        raise ValueError("score_not_numeric")
    value = float(value)
    if not math.isfinite(value) or not 0 <= value <= 1:
        raise ValueError("score_outside_unit_interval")
    return value


def branch_fields(row):
    if context_list(row["contexts"]):
        return "context", ("cross_score", "llm_score")
    return "abstract", ("abstract_cross_score", "abstract_llm_score")


def model_inputs(row):
    branch, _ = branch_fields(row)
    cited = clean_text(row["cited_abstract"])
    if not cited:
        raise MissingInput("missing_cited_abstract")
    if branch == "context":
        return branch, context_list(row["contexts"]), cited
    citing = clean_text(row["citing_abstract"])
    if not citing:
        raise MissingInput("missing_citing_abstract")
    return branch, [citing], cited





def load_seed_ids(path):
    from openpyxl import load_workbook
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        rows = workbook.active.iter_rows(values_only=True)
        header = next(rows)
        index = header.index("id")
        ids = []
        for row in rows:
            if not any(value is not None for value in row):
                continue
            value = row[index]
            if value is None or isinstance(value, bool):
                raise ValueError("missing_seed_id")
            parsed = int(value)
            if str(value).strip() not in (str(parsed), f"{parsed}.0") or parsed <= 0:
                raise ValueError("invalid_seed_id")
            ids.append(parsed)
        if len(ids) != 400 or len(set(ids)) != 400:
            raise ValueError("seed_file_requires_400_unique_ids")
        return sorted(ids)
    finally:
        workbook.close()


@contextmanager
def connect_database(read_only=True):
    """A dedicated SQLAlchemy Session; no transaction spans an LLM request.

    Read-only connections use DBAPI autocommit plus a PostgreSQL read-only
    default. The writer keeps one physical connection for the advisory lock,
    but each score update uses its own short Session.begin() transaction.
    NullPool closes that physical connection on exit, releasing the lock.
    """
    options = "-c statement_timeout=120000 -c lock_timeout=5000"
    if read_only:
        options += " -c default_transaction_read_only=on"
    url = os.getenv("DATABASE_URL")
    if url:
        from sqlalchemy.engine import make_url
        url = make_url(url).set(drivername="postgresql+psycopg")
    else:
        url = URL.create("postgresql+psycopg",
            host=os.getenv("POSTGRES_HOST", "localhost"),
            port=int(os.getenv("POSTGRES_PORT", "5432")),
            database=os.environ["POSTGRES_DB"], username=os.environ["POSTGRES_USER"],
            password=os.getenv("POSTGRES_PASS") or os.getenv("POSTGRES_PASSWORD"))
    engine = create_engine(url, poolclass=NullPool,
        isolation_level="AUTOCOMMIT" if read_only else "READ COMMITTED",
        connect_args={"connect_timeout": 10, "options": options})
    try:
        with engine.connect() as connection:
            with Session(bind=connection, expire_on_commit=False, autoflush=False) as session:
                yield session
    finally:
        engine.dispose()


def load_rows(conn, seed_ids):
    columns = {r["name"] for r in inspect(conn.connection()).get_columns(
        Citation.__tablename__, schema="public")}
    if set(FIELDS) - columns:
        raise StopRun("Database is missing required score columns; no automatic migration.")
    present = set(conn.scalars(select(Paper.corpus_id).where(
        Paper.corpus_id.in_(seed_ids))))
    if len(present) != 400:
        raise StopRun("Not all 400 seeds are present in the database.")
    return conn.execute(citation_statement().where(
        Citation.cited_corpus_id.in_(seed_ids)).order_by(
        Citation.cited_corpus_id, Citation.citing_corpus_id)).mappings().all()


def role_sets(authors):
    """Validate IDs and order using the same criteria as data_sampling.ipynb."""
    if not authors:
        raise MissingInput("missing_author_list_or_order")
    normalized = []
    for order, author in authors:
        for value in (order, author):
            if (isinstance(value, bool) or not isinstance(value, (int, float))
                    or not math.isfinite(value) or int(value) != value):
                raise MissingInput("invalid_author_id_or_order")
        if author <= 0:
            raise MissingInput("invalid_author_id_or_order")
        normalized.append((int(order), int(author)))
    orders, ids = zip(*normalized)
    first, last = min(orders), max(orders)
    if (len(set(ids)) != len(ids) or len(set(orders)) != len(orders)
            or first not in (0, 1)
            or sorted(orders) != list(range(first, first + len(orders)))):
        raise MissingInput("invalid_author_id_or_order")
    return (set(ids), {a for o, a in normalized if o == first},
            {a for o, a in normalized if o == last})


def eligible_input_rows(rows, graph):
    """Intersect collected text and author inputs, never score/date availability."""
    eligible, counts = [], Counter()
    counts["stored_citation_pairs"] = len(rows)
    # Keep explicit zeros in the audit for reproducible coverage comparisons.
    for key in ("text_ready_pairs", "valid_author_list_pairs", "eligible_pairs",
                "excluded_input_pairs", "eligible_context_pairs", "eligible_abstract_only_pairs"):
        counts[key] = 0
    for row in rows:
        authors_valid = all(row[key] in graph.roles for key in
                            ("citing_corpus_id", "cited_corpus_id"))
        counts["valid_author_list_pairs"] += int(authors_valid)
        try:
            branch, _, _ = model_inputs(row)
        except MissingInput as exc:
            counts[f"excluded_text_{exc}"] += 1
            continue
        counts["text_ready_pairs"] += 1
        if not authors_valid:
            counts["excluded_text_ready_invalid_author_lists"] += 1
            continue
        eligible.append(row)
        counts["eligible_pairs"] += 1
        counts["eligible_context_pairs" if branch == "context"
               else "eligible_abstract_only_pairs"] += 1
    counts["excluded_input_pairs"] = len(rows) - len(eligible)
    return eligible, dict(counts)


def author_score(citing, cited, neighbors):
    """0-hop overlap first, then 1-/2-hop links in the fixed local snapshot."""
    ca, cf, cl = citing
    da, df, dl = cited
    if cf & df:
        return 0.0
    if cl & dl:
        return 0.1
    if ca & da:
        return 0.5
    cn = set().union(*(neighbors(a) for a in ca))
    dn = set().union(*(neighbors(a) for a in da))
    if cn & da:
        return 0.6
    if cn & dn:
        return 0.8
    return 1.0


class AuthorGraph:
    def __init__(self, conn, rows):
        self.conn = conn
        self.roles = {}
        self.cache = {}
        paper_ids = sorted({r[k] for r in rows for k in
                            ("citing_corpus_id", "cited_corpus_id")})
        authors = defaultdict(list)
        for r in conn.execute(select(
            PaperAuthor.paper_id, PaperAuthor.author_id, PaperAuthor.author_order,
            Author.name.label("author_name")
        ).outerjoin(Author, Author.author_id == PaperAuthor.author_id)
            .where(PaperAuthor.paper_id.in_(paper_ids))).mappings():
            authors[r["paper_id"]].append(r)
        for paper_id, values in authors.items():
            if not all(clean_text(r["author_name"]) for r in values):
                continue
            try:
                self.roles[paper_id] = role_sets([
                    (r["author_order"], r["author_id"]) for r in values
                ])
            except MissingInput:
                # Leave just this paper unresolved, not the whole run.
                continue

    def load_neighbors(self, author_ids):
        missing = sorted(set(author_ids) - self.cache.keys())
        if not missing:
            return
        for author in missing:
            self.cache[author] = {}
        # The collection-year snapshot does not require full publication dates.
        # Unknown-year records are not silently attributed to the 2024 corpus.
        source, neighbor = aliased(PaperAuthor), aliased(PaperAuthor)
        query = (select(source.author_id, neighbor.author_id.label("neighbor_id"))
            .select_from(source)
            .join(neighbor, (neighbor.paper_id == source.paper_id) &
                           (neighbor.author_id != source.author_id))
            .join(Paper, Paper.corpus_id == source.paper_id)
            .where(source.author_id.in_(missing), neighbor.author_id > 0,
                   Paper.year == AUTHOR_GRAPH_YEAR)
            .distinct())
        # Do not keep a partly filled cache if the database query fails.
        try:
            found = self.conn.execute(query).mappings().all()
        except Exception:
            for author in missing:
                self.cache.pop(author, None)
            raise
        for r in found:
            self.cache[r["author_id"]][r["neighbor_id"]] = True

    def neighbors(self, author):
        return set(self.cache[author])

    def input_status(self, row):
        """Readiness of a NULL author score, independent of semantic inputs."""
        citing = self.roles.get(row["citing_corpus_id"])
        cited = self.roles.get(row["cited_corpus_id"])
        if not citing or not cited:
            return "missing_author_list_or_order"
        return "ready"

    def calculate(self, row):
        citing = self.roles.get(row["citing_corpus_id"])
        cited = self.roles.get(row["cited_corpus_id"])
        if not citing or not cited:
            raise MissingInput("missing_author_list_or_order")
        # Direct author overlap needs no date or co-authorship graph.
        if not (citing[0] & cited[0]):
            self.load_neighbors(citing[0] | cited[0])
        result = author_score(citing, cited, self.neighbors)
        meta = {"author_time_policy": AUTHOR_POLICY,
                "source_publication_year": AUTHOR_GRAPH_YEAR,
                "citation_specific_cutoff": False,
                "last_author_is_corresponding_proxy": True,
                "missing_full_dates": "allowed when stored publication year is 2024",
                "unknown_year_collaborations": "excluded",
                "author_roles_sha256": digest([sorted(s) for roles in (citing, cited) for s in roles]),
                "graph_sha256": digest({str(a): self.cache.get(a, {})
                                         for a in citing[0] | cited[0]})}
        return result, meta


class Journal:
    def __init__(self, path):
        self.path = path
        self.results = {}
        if path.exists():
            for line in path.read_text(encoding="utf-8").splitlines():
                try:
                    event = json.loads(line)
                    if event.get("event") == "result":
                        score_value(event["score"])
                        self.results[event["cache_key"]] = event
                except (ValueError, KeyError, TypeError):
                    # A truncated last line after interruption is not a result.
                    continue

    def append(self, event):
        self.path.parent.mkdir(parents=True, exist_ok=True)
        # Prefix a newline too, isolating a possible truncated final record.
        with self.path.open("a", encoding="utf-8") as handle:
            handle.write("\n" + json.dumps({"time": utc_now(), **event},
                                           ensure_ascii=False, default=str,
                                           allow_nan=False) + "\n")
            handle.flush()
            os.fsync(handle.fileno())
        if event.get("event") == "result":
            self.results[event["cache_key"]] = event


class Scorers:
    def __init__(self, args, journal):
        self.args, self.journal = args, journal
        self.cross = self.client = None
        self.calls = 0
        self.rng = random.Random(args.seed)

    def cross_score(self, texts, cited):
        import torch
        from sentence_transformers import CrossEncoder
        if self.cross is None:
            torch.manual_seed(self.args.seed)
            if torch.cuda.is_available():
                torch.cuda.manual_seed_all(self.args.seed)
            self.cross = CrossEncoder(CROSS_MODEL, device=self.args.device,
                                      max_length=512, revision=self.args.cross_revision,
                                      local_files_only=not self.args.allow_model_download)
        logits = self.cross.predict([(t, cited) for t in texts],
                                    activation_fn=torch.nn.Identity(),
                                    show_progress_bar=False)
        if not all(math.isfinite(float(x)) for x in logits):
            raise ValueError("nonfinite_cross_logits")
        scores = [1 / (1 + math.exp(-float(x))) if float(x) >= 0 else
                  math.exp(float(x)) / (1 + math.exp(float(x))) for x in logits]
        revision = getattr(self.cross.model.config, "_commit_hash", None)
        return score_value(max(scores)), {"model": CROSS_MODEL,
                "resolved_revision": revision, "activation": "sigmoid_once",
                "aggregation": "max", "max_length": 512}

    def llm_score(self, prompt, ids):
        from openai import OpenAI
        if self.client is None:
            key = os.getenv("LLM_KEY") or os.getenv("OPENAI_API_KEY")
            if not key:
                raise StopRun("LLM_KEY or OPENAI_API_KEY is required for LLM scoring.")
            self.client = OpenAI(api_key=key, base_url=os.getenv("LLM_URL") or None,
                                 timeout=self.args.timeout, max_retries=0)
        settings = {"model": self.args.model, "response_format": {"type": "json_object"}}
        if self.args.temperature != "omit":
            settings["temperature"] = float(self.args.temperature)
        for attempt in range(self.args.retries):
            if self.calls >= self.args.max_llm_calls:
                raise StopRun("LLM request budget reached; re-run to resume.")
            self.calls += 1
            try:
                response = self.client.chat.completions.create(
                    messages=[{"role": "system", "content": SYSTEM_PROMPT},
                              {"role": "user", "content": prompt}], **settings)
                raw = response.choices[0].message.content or ""
                payload = json.loads(raw)
                value = score_value(payload["score"])
                usage = response.usage.model_dump() if response.usage else None
                return value, {"requested_model": self.args.model,
                    "returned_model": response.model, "request_id": response.id,
                    "response": payload, "usage": usage, "settings": settings,
                    "attempt": attempt + 1, "prompt_sha256": digest(prompt)}
            except Exception as exc:
                status = getattr(exc, "status_code", None)
                self.journal.append({"event": "llm_attempt_failed", **ids,
                    "attempt": attempt + 1, "error_type": type(exc).__name__,
                    "http_status": status})
                if status in (400, 401, 403, 404):
                    raise StopRun(f"LLM endpoint rejected request (HTTP {status}); check model/settings.") from None
                if attempt + 1 == self.args.retries:
                    raise MissingInput("llm_retries_exhausted") from None
                time.sleep(min(30, 2 ** (attempt + 1)) + self.rng.random())
            finally:
                if self.args.interval:
                    time.sleep(self.args.interval)


def snapshot_hash(row):
    return digest({k: row[k] for k in ("contexts", "citing_abstract", "cited_abstract",
                                     "citing_date", "citing_year")})


def write_result(conn, row, field, value, replace_author=False):
    """NULL-only semantic writes; opt-in author replacement with input/value guards."""
    if field not in FIELDS:
        raise ValueError("unknown_score_field")
    replacing = field == "author_independence" and replace_author
    if row[field] is not None and not replacing:
        return False
    keys = (Citation.citing_corpus_id == row["citing_corpus_id"],
            Citation.cited_corpus_id == row["cited_corpus_id"])
    with conn.begin():
        current = conn.execute(citation_statement().where(*keys)
            .with_for_update(of=(Citation, CitingPaper, CitedPaper))).mappings().first()
        if current is None or snapshot_hash(current) != snapshot_hash(row):
            raise MissingInput("input_changed_before_commit")
        # Replacement is opt-in, author-only and compare-and-swap protected.
        expected_value = (getattr(Citation, field).is_(None) if row[field] is None
                          else getattr(Citation, field) == row[field])
        result = conn.execute(update(Citation)
            .where(*keys, expected_value)
            .values({field: score_value(value), "scored_at": func.now()})
            .execution_options(synchronize_session=False))
        return result.rowcount == 1


def audit_rows(rows, graph, recompute_author=False):
    records, counts = [], Counter()
    for row in rows:
        record = {k: row[k] for k in ("citing_corpus_id", "cited_corpus_id")}
        # This is a report flag, NOT a model field or a substitute for writing U.
        # Author readiness must be reported even if context parsing fails.
        needs_author = row["author_independence"] is None or recompute_author
        record["needs_author_independence"] = needs_author
        record["author_time_policy"] = AUTHOR_POLICY
        if row["author_independence"] is None:
            counts["missing_author_independence"] += 1
        elif recompute_author:
            counts["recompute_existing_author_independence"] += 1
        if needs_author:
            counts["pending_author_independence"] += 1
            record["author_status"] = graph.input_status(row)
            counts[f"author_{record['author_status']}"] += 1
        else:
            record["author_status"] = "existing_preserved"
        try:
            branch, fields = branch_fields(row)
            counts[branch] += 1
            record["branch"] = branch
            missing = [f for f in fields if row[f] is None]
            record["missing_semantic_fields"] = ",".join(missing)
            for field in missing:
                counts[f"missing_{field}"] += 1
            try:
                model_inputs(row)
                record["input_status"] = "ready"
                for field in missing:
                    counts[f"ready_{field}"] += 1
            except MissingInput as exc:
                record["input_status"] = str(exc)
                counts[str(exc)] += 1
            for field in (*fields, "author_independence"):
                if row[field] is not None:
                    try:
                        score_value(row[field])
                    except ValueError:
                        counts[f"invalid_existing_{field}"] += 1
        except MissingInput as exc:
            record.update(branch="unresolved", input_status=str(exc))
            counts[str(exc)] += 1
        records.append(record)
    return records, dict(counts)


def pending_targets(row, component="all", branch_filter="all", recompute_author=False):
    """Use NULL fields, or explicit author refresh, never audit flags, to select work.

    Caller must first restrict rows to the text-ready, author-valid cohort.
    Within that cohort U does not require an existing semantic model result.
    The branch filter limits semantic scoring only, never author scoring.
    Explicit component restrictions still limit the requested run.
    """
    if component not in ("all", "author", "cross", "llm"):
        raise ValueError("unknown_scoring_component")
    if branch_filter not in ("all", "context", "abstract"):
        raise ValueError("unknown_branch_filter")
    try:
        branch, fields = branch_fields(row)
    except MissingInput:
        branch, fields = "unresolved", ()
    targets = []
    if component in ("all", "author") and (row["author_independence"] is None or recompute_author):
        targets.append("author_independence")
    if branch_filter != "all" and branch != branch_filter:
        return tuple(targets)
    if component != "author":
        targets.extend(field for field in fields
                       if row[field] is None and (component == "all" or component in field))
    return tuple(targets)


def scoring_tasks(selected):
    """Attempt all selected author fields before any potentially paid work."""
    for row, targets in selected:
        if "author_independence" in targets:
            yield row, "author_independence"
    for row, targets in selected:
        for field in targets:
            if field != "author_independence" and row[field] is None:
                yield row, field


def run(args):
    from dotenv import load_dotenv
    load_dotenv(BASE / ".env", override=False)
    load_dotenv(BASE.parent / ".env", override=False)
    try:
        args.seed = int(os.environ["SEED"])
        if args.seed < 0:
            raise ValueError
    except (KeyError, TypeError, ValueError):
        raise StopRun("SEED must be set to a nonnegative integer; no default is supplied.") from None
    endpoint = os.getenv("LLM_URL") or "https://api.openai.com/v1"
    provider_host = urlsplit(endpoint).hostname
    endpoint_signature = digest(endpoint)
    seed_ids = load_seed_ids(args.seeds)
    # Separate reader/writer connections: never hold a transaction during a call.
    with connect_database() as reader:
        stored_rows = load_rows(reader, seed_ids)
        graph = AuthorGraph(reader, stored_rows)
        rows, cohort_counts = eligible_input_rows(stored_rows, graph)
        audit, counts = audit_rows(rows, graph, args.recompute_author)
        args.output.mkdir(parents=True, exist_ok=True)
        with (args.output / "coverage.csv").open("w", newline="", encoding="utf-8") as handle:
            names = sorted({k for r in audit for k in r})
            writer = csv.DictWriter(handle, fieldnames=names)
            writer.writeheader()
            writer.writerows(audit)
        settings = {"time": utc_now(), "seed_count": len(seed_ids),
            "seed_ids_sha256": digest(seed_ids), "citation_pairs": len(rows),
            "pair_ids_sha256": digest([(r["citing_corpus_id"], r["cited_corpus_id"]) for r in rows]),
            "input_cohort": "Required text inputs intersect valid author lists on both papers",
            "date_eligibility_filter": False,
            "cohort_counts": cohort_counts,
            "author_time_policy": AUTHOR_POLICY,
            "author_graph_publication_year": AUTHOR_GRAPH_YEAR,
            "author_recompute_existing": args.recompute_author,
            "counts": counts, "base_seed": args.seed, "model": args.model,
            "provider_host": provider_host, "endpoint_sha256": endpoint_signature,
            "cross_model": CROSS_MODEL, "cross_revision": args.cross_revision,
            "temperature": args.temperature, "retry_limit": args.retries,
            "context_prompt_version": CONTEXT_PROMPT_VERSION,
            "abstract_prompt_version": ABSTRACT_PROMPT_VERSION,
            "max_llm_calls": args.max_llm_calls, "overwrite_semantic_scores": False,
            "script_sha256": hashlib.sha256(Path(__file__).read_bytes()).hexdigest()}
        (args.output / "audit.json").write_text(json.dumps(settings, indent=2), encoding="utf-8")
        print(json.dumps({"seed_count": 400, **cohort_counts,
                          "citation_pairs": len(rows), **counts}, indent=2))
        if not args.recompute_author and any(r["author_independence"] is not None for r in rows):
            print("WARNING: Existing author scores are preserved and their policy is unverified. "
                  "Use --recompute-author to obtain a consistent 2024-snapshot cohort.", flush=True)
        if not args.write:
            print("Read-only audit: no database changes, no model calls.")
            return 0
        selected = []
        for row in rows:
            targets = pending_targets(row, args.component, args.branch, args.recompute_author)
            if targets:
                selected.append((row, targets))
        if args.limit:
            # A small trial should include computable work, not merely the
            # first records with missing metadata. Ordering otherwise stays fixed.
            audit_by_pair = {(r["citing_corpus_id"], r["cited_corpus_id"]): r for r in audit}
            def computable(item):
                row, targets = item
                info = audit_by_pair[(row["citing_corpus_id"], row["cited_corpus_id"])]
                return any((row[f] is None or (f == "author_independence" and args.recompute_author)) and
                    info.get("author_status" if f == "author_independence" else "input_status") == "ready"
                    for f in targets)
            selected.sort(key=lambda item: not computable(item))
            selected = selected[:args.limit]
        journal = Journal(args.output / "events.jsonl")
        journal.append({"event": "run_start", **settings, "selected_pairs": len(selected),
                        "branch_filter": args.branch, "component": args.component})
        scorers, totals = Scorers(args, journal), Counter()
        print(f"Selected pending citation pairs: {len(selected)}; LLM attempt cap: {args.max_llm_calls}")
        tasks = list(scoring_tasks(selected))
        author_tasks = sum(field == "author_independence" for _, field in tasks)
        print(f"Author pass: {author_tasks} fields; policy={AUTHOR_POLICY}; "
              f"replace existing={args.recompute_author}.")
        with connect_database(read_only=False) as writer:
            # Prevent concurrent runs for this seed cohort from duplicating
            # paid calls. The session lock is released when the connection closes.
            lock_id = int(digest(["score_all_citations", seed_ids])[:15], 16)
            locked = writer.scalar(select(func.pg_try_advisory_lock(lock_id)))
            writer.commit()
            if not locked:
                raise StopRun("Another scoring run holds the lock for this seed cohort.")
            for index, (row, field) in enumerate(tasks, 1):
                ids = {k: row[k] for k in ("citing_corpus_id", "cited_corpus_id")}
                # One field per task; author work for the whole selection comes first.
                if row[field] is not None and not (field == "author_independence" and args.recompute_author):
                    continue
                try:
                    if field == "author_independence":
                        value, meta = graph.calculate(row)
                        event = {"event": "result", **ids, "field": field,
                                 "score": value, "metadata": meta,
                                 "previous_score": row[field] if row[field] is None or
                                     (isinstance(row[field], (int, float)) and math.isfinite(row[field]))
                                     else str(row[field]),
                                 "cache_key": digest([ids, field, snapshot_hash(row), meta])}
                        journal.append(event)
                    else:
                        branch, texts, cited = model_inputs(row)
                        prompt = build_prompt(branch, texts, cited)
                        signature = {"field": field, "ids": ids, "texts": texts, "cited": cited,
                            "model": args.model if "llm" in field else CROSS_MODEL,
                            "endpoint_sha256": endpoint_signature if "llm" in field else None,
                            "revision": args.cross_revision, "max_length": 512,
                            "temperature": args.temperature, "prompt": prompt,
                            "system": SYSTEM_PROMPT, "pipeline": "raw_scores_v1"}
                        key = digest(signature)
                        event = journal.results.get(key)
                        if event:
                            value = event["score"]
                            totals["reused_checkpoint_results"] += 1
                        else:
                            value, meta = (scorers.llm_score(prompt, ids) if "llm" in field
                                           else scorers.cross_score(texts, cited))
                            event = {"event": "result", **ids, "field": field, "score": value,
                                     "metadata": meta, "cache_key": key,
                                     "input_sha256": digest([branch, texts, cited])}
                            journal.append(event)
                    changed = write_result(writer, row, field, value, args.recompute_author)
                    journal.append({"event": "committed" if changed else "existing_preserved",
                                    **ids, "field": field, "score": value,
                                    "cache_key": event["cache_key"]})
                    totals[f"written_{field}" if changed else "concurrent_existing_preserved"] += 1
                except MissingInput as exc:
                    totals[str(exc)] += 1
                    if field == "author_independence":
                        totals["unresolved_author_independence"] += 1
                    journal.append({"event": "unresolved", **ids, "field": field,
                                    "reason": str(exc)})
                except StopRun:
                    raise
                except Exception as exc:
                    # Do not print/log exception text (may contain DB URL or API key).
                    journal.append({"event": "failed", **ids, "field": field,
                                    "error_type": type(exc).__name__})
                    raise StopRun(f"Stopped on {type(exc).__name__}; completed fields are saved.") from None
                if index == author_tasks and author_tasks:
                    print(f"Author pass finished: attempted={author_tasks}; "
                          f"written={totals['written_author_independence']}; "
                          f"unresolved={totals['unresolved_author_independence']}; "
                          f"existing preserved={totals['concurrent_existing_preserved']}", flush=True)
                if index % 10 == 0 or index == len(tasks):
                    print(f"Processed fields {index}/{len(tasks)}; API attempts={scorers.calls}; {dict(totals)}", flush=True)
        journal.append({"event": "run_end", "totals": dict(totals), "api_attempts": scorers.calls})
        print("Run finished. Re-run without --write to audit remaining fields.")
        return 0


def self_test():
    """Offline boundary tests, no credentials/database/model required."""
    assert context_list(None) == context_list([]) == context_list([" "]) == []
    assert context_list([{"text": "abc"}, " def "]) == ["abc", "def"]
    assert context_list(["[1]"]) == ["[1]"]
    for bad in (True, "0.2", float("nan"), float("inf"), -0.1, 1.1):
        try:
            score_value(bad)
        except ValueError:
            pass
        else:
            raise AssertionError("invalid_score_accepted")
    assert score_value(0) == 0 and score_value(1) == 1
    row = {"contexts": ["context"], "cited_abstract": "cited", "citing_abstract": None}
    assert branch_fields(row)[1] == ("cross_score", "llm_score")
    assert model_inputs(row) == ("context", ["context"], "cited")
    row.update(contexts=[], citing_abstract="citing")
    assert branch_fields(row)[1] == ("abstract_cross_score", "abstract_llm_score")
    assert model_inputs(row) == ("abstract", ["citing"], "cited")
    row["citing_abstract"] = None
    try:
        model_inputs(row)
    except MissingInput:
        pass
    else:
        raise AssertionError("missing_abstract_must_remain_unresolved")
    pending = {
        "contexts": None, "cited_abstract": None, "citing_abstract": None,
        "cross_score": None, "llm_score": None,
        "abstract_cross_score": None, "abstract_llm_score": None,
        "author_independence": None,
    }
    assert pending_targets(pending, "author") == ("author_independence",)
    # Task selection itself does not use semantic-score availability. The
    # collected-input eligibility filter runs before this helper in run().
    assert pending_targets(pending, "all", "context") == ("author_independence",)
    assert pending_targets(pending) == (
        "author_independence", "abstract_cross_score", "abstract_llm_score")
    pending["contexts"] = ["context"]
    assert pending_targets(pending, "all", "abstract") == ("author_independence",)
    pending["contexts"] = {"unsupported": "context"}
    assert pending_targets(pending) == ("author_independence",)
    pending["author_independence"] = 0.0
    assert pending_targets(pending, "author") == ()
    assert pending_targets(pending) == ()
    # All author tasks precede semantic tasks, so an LLM budget stop cannot
    # leave the remainder of the selected author pass unattempted.
    first = dict(pending, contexts=["context"], author_independence=None)
    second = dict(pending, contexts=[], author_independence=None)
    tasks = list(scoring_tasks([(r, pending_targets(r)) for r in (first, second)]))
    assert [field for _, field in tasks[:2]] == ["author_independence"] * 2
    assert all(field != "author_independence" for _, field in tasks[2:])
    # Author validity matches the notebook, including both 0-/1-based order.
    assert role_sets([(0, 1), (1, 2)]) == role_sets([(1, 1), (2, 2)])
    for invalid in ([], [(None, 1)], [(1, 0)], [(True, 1)], [(1, True)],
                    [(1, 1), (3, 2)], [(1, 1), (1, 2)], [(1, 1), (2, 1)],
                    [(1.5, 1)], [(2, 1)]):
        try:
            role_sets(invalid)
        except MissingInput:
            pass
        else:
            raise AssertionError("invalid_author_list_accepted")
    a = role_sets([(1, 1), (2, 2)])
    assert author_score(a, a, None) == 0
    assert author_score(a, role_sets([(1, 3), (2, 2)]), None) == 0.1
    assert author_score(role_sets([(1, 1), (2, 2), (3, 3)]),
                        role_sets([(1, 4), (2, 2), (3, 5)]), None) == 0.5
    b = role_sets([(1, 3), (2, 4)])
    def neighbors(author):
        return {1: {3}, 3: {1}}.get(author, set())
    assert author_score(a, b, neighbors) == 0.6
    assert author_score(a, b, lambda _: {5}) == 0.8
    assert author_score(a, b, lambda _: set()) == 1
    # No citing date is needed under the expressly chosen snapshot policy.
    graph = AuthorGraph.__new__(AuthorGraph)
    graph.roles, graph.cache = {11: a, 100: b}, {1: {3: True}, 2: {}, 3: {1: True}, 4: {}}
    graph_row = dict(first, citing_corpus_id=11, cited_corpus_id=100,
                     citing_date=None, citing_year=None)
    assert graph.input_status(graph_row) == "ready"
    result, meta = graph.calculate(graph_row)
    assert result == 0.6 and meta["citation_specific_cutoff"] is False
    assert meta["source_publication_year"] == 2024
    # Dates and saved scores never filter the collected-input intersection.
    candidates = [
        dict(graph_row, cited_abstract="cited", citing_abstract=None),
        dict(graph_row, contexts=[], cited_abstract="cited", citing_abstract="citing"),
        dict(graph_row, cited_abstract=None),
        dict(graph_row, citing_corpus_id=12, cited_abstract="cited"),
    ]
    eligible, counts = eligible_input_rows(candidates, graph)
    assert eligible == candidates[:2]
    assert counts["text_ready_pairs"] == 3 and counts["valid_author_list_pairs"] == 3
    assert counts["eligible_pairs"] == 2
    # Explicit refresh changes U only: semantic zeros and other values stay intact.
    complete = dict(graph_row, cross_score=0.0, llm_score=0.7, author_independence=1.0)
    assert pending_targets(complete) == ()
    assert pending_targets(complete, recompute_author=True) == ("author_independence",)
    refreshed = list(scoring_tasks([(complete, pending_targets(complete, recompute_author=True))]))
    assert refreshed == [(complete, "author_independence")]
    assert write_result(None, complete, "cross_score", 0.4, replace_author=True) is False
    assert write_result(None, complete, "author_independence", 0.8) is False
    # Exercise replacement SQL without opening a database or executing writes.
    from contextlib import nullcontext
    class FakeResult:
        rowcount = 1
        def mappings(self):
            return self
        def first(self):
            return complete
    class FakeWriter:
        def __init__(self):
            self.statements = []
        def begin(self):
            return nullcontext()
        def execute(self, statement):
            self.statements.append(statement)
            return FakeResult()
    writer = FakeWriter()
    assert write_result(writer, complete, "author_independence", 0.6, replace_author=True)
    params = writer.statements[-1].compile().params
    assert params["author_independence"] == 0.6
    assert params["author_independence_1"] == 1.0
    assert len(writer.statements) == 2
    print("Offline tests passed: text/author cohort, date-independent snapshot, 0-/1-/2-hop "
          "mapping, NULL-only semantics, explicit author refresh and compare-and-swap protection.")


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--seeds", type=Path, default=BASE / "sample/seed_papers_final.xlsx")
    parser.add_argument("--output", type=Path, default=BASE / "sample/citation_scoring")
    parser.add_argument("--write", action="store_true",
                        help="Enable missing-score writes and model calls; existing author scores "
                             "change only with --recompute-author")
    parser.add_argument("--limit", type=int, default=0, help="Maximum pending citation pairs; 0 selects all")
    parser.add_argument("--branch", choices=("all", "context", "abstract"), default="all",
                        help="Semantic branch filter; NULL author scores are attempted on both branches")
    parser.add_argument("--component", choices=("all", "cross", "llm", "author"), default="all")
    parser.add_argument("--recompute-author", action="store_true",
                        help="Recompute all eligible author scores under the 2024-snapshot policy; "
                             "journal old values, preserve semantic scores")
    parser.add_argument("--model", default="gpt-5-mini")
    parser.add_argument("--temperature", default="0", help="Project setting 0, or omit if endpoint requires it")
    parser.add_argument("--cross-revision", default=CROSS_REVISION, help="Frozen Hugging Face commit hash")
    parser.add_argument("--device", default="cpu")
    parser.add_argument("--allow-model-download", action="store_true")
    parser.add_argument("--max-llm-calls", type=int, default=10, help="Hard cap including retries (default 10)")
    parser.add_argument("--retries", type=int, default=3)
    parser.add_argument("--timeout", type=float, default=120)
    parser.add_argument("--interval", type=float, default=0)
    parser.add_argument("--self-test", action="store_true")
    args = parser.parse_args()
    if args.recompute_author and args.component not in ("all", "author"):
        parser.error("--recompute-author requires --component all or author")
    if args.limit < 0 or args.max_llm_calls < 0 or args.retries < 1 or args.timeout <= 0 or args.interval < 0:
        parser.error("Invalid limit/budget/retry/timeout/interval")
    if args.temperature != "omit":
        try:
            temperature = float(args.temperature)
            if not math.isfinite(temperature) or not 0 <= temperature <= 2:
                raise ValueError
        except ValueError:
            parser.error("temperature must be omit or a number in [0,2]")
    if args.self_test:
        self_test()
        return 0
    try:
        return run(args)
    except KeyboardInterrupt:
        print("Interrupted; committed scores and checkpointed results are retained.", file=sys.stderr)
        return 130
    except StopRun as exc:
        print(str(exc), file=sys.stderr)
        return 2
    except Exception as exc:
        print(f"Stopped on {type(exc).__name__}; no credentials/exception details are logged.", file=sys.stderr)
        return 2


if __name__ == "__main__":
    sys.exit(main())
